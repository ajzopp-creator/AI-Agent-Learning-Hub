"""Permanent tests for WO-P025-EN.003 Positions-by-account formulas."""

from __future__ import annotations

from openpyxl import Workbook

from domain.formula_templates import (
    positions_cost_formula,
    positions_headers,
    positions_shares_formula,
)
from infrastructure.analytics_sheets import _fifo_cost_pairs, build_positions


def test_cost_formula_is_sumifs_by_ticker_and_account():
    f = positions_cost_formula(3)
    assert "SUMIFS" in f
    assert "Fifo_Cost" in f
    assert "A3" in f
    assert "I3" in f


def test_shares_formula_is_sumifs_fifo_qty():
    f = positions_shares_formula(4)
    assert "SUMIFS" in f
    assert "Fifo_Cost!$C:$C" in f
    assert "I4" in f


def test_headers_include_account():
    assert positions_headers()[-1] == "Account"


def test_fifo_cost_pairs_unique_and_capped():
    wb = Workbook()
    ws = wb.active
    ws.title = "Fifo_Cost"
    ws.append(["Ticker", "Account", "Remaining_Shares", "Remaining_Cost"])
    ws.append(["SPY", "AJZ6348", 10, 4000])
    ws.append(["SPY", "5232-9885", 3, 1230])
    ws.append(["SPY", "AJZ6348", 1, 400])
    pairs = _fifo_cost_pairs(wb, cap=200)
    assert pairs == [("SPY", "AJZ6348"), ("SPY", "5232-9885")]


def test_build_positions_writes_account_not_primary():
    wb = Workbook()
    wb.active.title = "Positions"
    fc = wb.create_sheet("Fifo_Cost")
    fc.append(["Ticker", "Account", "Remaining_Shares", "Remaining_Cost"])
    fc.append(["AAPL", "AJZ6348", 8, 1600])
    fc.append(["AAPL", "5232-9885", 2, 400])
    wb.create_sheet("Market_Data")
    build_positions(wb)
    pos = wb["Positions"]
    assert "by account" in str(pos.cell(1, 1).value).lower()
    assert pos.cell(3, 1).value == "AAPL"
    assert pos.cell(3, 9).value == "AJZ6348"
    assert pos.cell(4, 9).value == "5232-9885"
    assert "SUMIFS" in str(pos.cell(3, 5).value)
    assert "PRIMARY" not in {pos.cell(3, 9).value, pos.cell(4, 9).value}
