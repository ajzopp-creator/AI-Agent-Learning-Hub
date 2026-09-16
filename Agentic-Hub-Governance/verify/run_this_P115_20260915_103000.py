"""
run_this_P115_20260915_103000.py
Purpose: Pull 27-col tracker rows Date >= 2026-09-01 for the WO-P115-E6.001
Phase 1 Z-Score Regime Council pilot closing summary. Read-only.
Prints raw Comments field (where Z-scores are logged per ways-of-working.md
convention) alongside Symbol/Step1Verdict/Date so Claude can parse and
tabulate the regime distribution vs actual verdicts without guessing the
shorthand format blind.
"""
import sys
from datetime import datetime
import openpyxl

sys.stdout.reconfigure(encoding='utf-8', errors='replace')

TRACKER_PATH = r"D:\OneDrive\Documents\AJZStrategiesLLC\P_115_TrackerAudit\P_115_118_TrackerDashboard_V2.xlsx"
CUTOFF = datetime(2026, 9, 1)

def main():
    wb = openpyxl.load_workbook(TRACKER_PATH, data_only=True)
    print("SHEETS:", wb.sheetnames)
    ws = wb['Tracker Log'] if 'Tracker Log' in wb.sheetnames else wb[wb.sheetnames[0]]

    header_row = [c.value for c in ws[1]]
    print("HEADERS:", header_row)

    def col_idx(name):
        for i, h in enumerate(header_row):
            if h and str(h).strip().lower() == name.strip().lower():
                return i
        return None

    idx_date = col_idx("Date")
    idx_symbol = col_idx("Symbol")
    idx_source = col_idx("SignalSource")
    idx_verdict = col_idx("Step1Verdict")
    idx_comments = col_idx("Comments")
    idx_outcome = col_idx("Outcome")

    print("COL_IDX date=%s symbol=%s source=%s verdict=%s comments=%s outcome=%s" % (
        idx_date, idx_symbol, idx_source, idx_verdict, idx_comments, idx_outcome))

    if None in (idx_date, idx_symbol, idx_verdict, idx_comments):
        print("MISSING_REQUIRED_COLUMNS -- aborting")
        sys.exit(1)

    rows_out = []
    for r in ws.iter_rows(min_row=2, values_only=True):
        raw_date = r[idx_date]
        if raw_date is None:
            continue
        if isinstance(raw_date, datetime):
            dt = raw_date
        else:
            try:
                dt = datetime.strptime(str(raw_date).strip(), "%m/%d/%Y")
            except ValueError:
                try:
                    dt = datetime.strptime(str(raw_date).strip(), "%m/%d/%y")
                except ValueError:
                    continue
        if dt < CUTOFF:
            continue
        symbol = r[idx_symbol]
        source = r[idx_source] if idx_source is not None else None
        verdict = r[idx_verdict]
        comments = r[idx_comments]
        outcome = r[idx_outcome] if idx_outcome is not None else None
        rows_out.append((dt.strftime("%Y-%m-%d"), symbol, source, verdict, outcome, comments))

    print("ROW_COUNT:", len(rows_out))
    print("---ROWS---")
    for row in rows_out:
        print("|".join("" if v is None else str(v) for v in row))

if __name__ == "__main__":
    main()
