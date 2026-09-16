# C:\Users\Trader\AI-Agent-Learning-Hub\Agentic-Hub-Governance\verify\run_this_P025_20260915_215200.py
"""
PEH verify-only — WO-P025-EN.001 Independent Review workbook + gate checks.

Does NOT build, format, copy, rmtree, or mark the WO CLOSED.
Run from Agentic-Hub-Governance\\verify\\ with HUB_ROOT set.
"""

from __future__ import annotations

import os
import sys
from datetime import datetime
from pathlib import Path

HUB_ROOT = Path(os.environ.get("HUB_ROOT", r"C:\Users\Trader\AI-Agent-Learning-Hub"))
PROJECT_ROOT = HUB_ROOT / "projects" / "P_025_AJZ_Institutional_Portfolio_Tracker"
PYTHON_ROOT = PROJECT_ROOT / "python"
OUTPUT_DIR = PROJECT_ROOT / "output"
LAKE_PATH = OUTPUT_DIR / "P_025_Portfolio_BUILT.xlsx"
WO_PATH = HUB_ROOT / "Agentic-Hub-Governance" / "work_orders" / "WO-P025-EN.001.md"


def write_done(status: str, exit_code: int) -> None:
    Path(__file__).with_suffix(Path(__file__).suffix + ".done").write_text(
        f"status={status}\nexit_code={exit_code}\ntimestamp={datetime.now().isoformat()}\n",
        encoding="utf-8",
    )


def fail(msg: str) -> None:
    print(f"FAIL: {msg}")
    write_done("FAIL", 1)
    sys.exit(1)


def latest_analytics() -> Path:
    files = sorted(
        OUTPUT_DIR.glob("P_025_Portfolio_BUILT_Analytics_*.xlsx"),
        key=lambda p: p.stat().st_mtime,
        reverse=True,
    )
    if not files:
        fail(f"No Analytics workbook under {OUTPUT_DIR}")
    return files[0]


def header_map(ws) -> dict[str, int]:
    return {
        str(ws.cell(1, c).value).strip().lower(): c
        for c in range(1, (ws.max_column or 1) + 1)
        if ws.cell(1, c).value
    }


def main() -> int:
    print("=" * 60)
    print("P_025 PEH 20260915_215200  VERIFY ONLY  WO-P025-EN.001")
    print("=" * 60)
    print(f"HUB_ROOT : {HUB_ROOT}")
    print("Mode     : no build / no format / no CLOSE")

    if "OneDrive" in str(HUB_ROOT) and not os.environ.get("HUB_ROOT"):
        fail("HUB_ROOT resolved to OneDrive without an explicit env override")
    if not PYTHON_ROOT.exists():
        fail(f"python\\ missing: {PYTHON_ROOT}")
    if str(PYTHON_ROOT) not in sys.path:
        sys.path.insert(0, str(PYTHON_ROOT))

    import config as cfg

    if "OneDrive" in cfg._DEFAULT_HUB_ROOT:
        fail("config._DEFAULT_HUB_ROOT still contains OneDrive")
    print(f"  OK  default hub = {cfg._DEFAULT_HUB_ROOT}")
    print(f"  OK  IRA_FEED_READY = {cfg.IRA_FEED_READY}")

    if not WO_PATH.exists():
        fail(f"WO not on Hub ledger: {WO_PATH}")
    print(f"  OK  WO ledger file {WO_PATH.name}")

    print("\n=== pytest (read-only) ===")
    import pytest

    rc = pytest.main(["-q", str(PYTHON_ROOT / "tests")])
    if rc != 0:
        fail(f"pytest failed with code {rc}")
    print("  OK  pytest")

    if not LAKE_PATH.exists():
        fail(f"Data Lake missing: {LAKE_PATH}")
    analytics = latest_analytics()
    print(f"  OK  lake      {LAKE_PATH.name} ({LAKE_PATH.stat().st_size / 1024:.1f} KB)")
    print(f"  OK  analytics {analytics.name} ({analytics.stat().st_size / 1024:.1f} KB)")

    from openpyxl import load_workbook

    print("\n=== 1–2 Data Lake: Fifo sheets + Cost_Basis label ===")
    lake = load_workbook(LAKE_PATH, read_only=True, data_only=False)
    for sheet in ("Fifo_Lots", "Fifo_Cost", "Cost_Basis", "Trade_Log"):
        if sheet not in lake.sheetnames:
            fail(f"Lake missing sheet {sheet}")
    print("  OK  Fifo_Lots + Fifo_Cost present")

    cb_headers = [str(lake["Cost_Basis"].cell(1, c).value or "") for c in range(1, 6)]
    joined = " ".join(cb_headers)
    if "fifo" in joined.lower():
        fail(f"Cost_Basis header names FIFO: {cb_headers}")
    print(f"  OK  Cost_Basis headers {cb_headers}")

    fifo_n = max(lake["Fifo_Lots"].max_row - 1, 0)
    cost_n = max(lake["Fifo_Cost"].max_row - 1, 0)
    print(f"  OK  fifo_lots rows={fifo_n}  fifo_cost rows={cost_n}")

    print("\n=== spot-check Fifo_Lots vs open longs ===")
    tl = lake["Trade_Log"]
    cols = header_map(tl)
    need = ("underlying_symbol", "direction", "status")
    if not all(k in cols for k in need):
        print(f"  WARN Trade_Log headers missing {need}; skip open-long compare")
        open_longs: set[str] = set()
    else:
        open_longs = set()
        for r in range(2, (tl.max_row or 1) + 1):
            direction = str(tl.cell(r, cols["direction"]).value or "").strip().lower()
            status = str(tl.cell(r, cols["status"]).value or "").strip().lower()
            sym = str(tl.cell(r, cols["underlying_symbol"]).value or "").strip().upper()
            if sym and direction == "long" and status == "open":
                open_longs.add(sym)
        lot_tickers = set()
        lots = lake["Fifo_Lots"]
        for r in range(2, (lots.max_row or 1) + 1):
            t = lots.cell(r, 2).value
            if t:
                lot_tickers.add(str(t).strip().upper())
        extra = sorted(lot_tickers - open_longs)
        missing = sorted(open_longs - lot_tickers)
        print(f"  OK  open longs={len(open_longs)}  fifo tickers={len(lot_tickers)}")
        if extra:
            print(f"  WARN fifo tickers not in open-long set: {extra[:12]}")
        if missing:
            print(f"  WARN open longs with no fifo lot: {missing[:12]}")
        if not extra and not missing:
            print("  OK  fifo tickers match open-long set")
    lake.close()

    print("\n=== 3–6 Analytics: Positions / CORREL / Geo / Stress ===")
    wb = load_workbook(analytics, read_only=True, data_only=False)
    for sheet in ("Positions", "Correlation", "Geographic_Exposure", "Stress_Testing"):
        if sheet not in wb.sheetnames:
            fail(f"Analytics missing sheet {sheet}")

    pos_cost = wb["Positions"].cell(3, 5).value or ""
    if "SUMIF" not in str(pos_cost) or "Fifo_Cost" not in str(pos_cost):
        fail(f"Positions!E3 is not Fifo_Cost SUMIF: {pos_cost!r}")
    print(f"  OK  Positions!E3 {pos_cost}")

    sample = wb["Correlation"].cell(5, 3).value
    if not (isinstance(sample, str) and "CORREL" in sample):
        fail(f"Correlation!C5 is not CORREL: {sample!r}")
    print("  OK  Correlation off-diagonal CORREL")

    geo = str(wb["Geographic_Exposure"].cell(1, 1).value or "")
    stress = str(wb["Stress_Testing"].cell(1, 1).value or "")
    if "placeholder" in geo.lower():
        fail(f"Geographic still placeholder: {geo!r}")
    if "placeholder" in stress.lower():
        fail(f"Stress still placeholder: {stress!r}")
    print(f"  OK  Geographic title={geo}")
    print(f"  OK  Stress title={stress}")
    wb.close()

    print("\n" + "=" * 60)
    print("PASS — verify only. WO-P025-EN.001 stays OPEN until CLOSE stamp.")
    print("=" * 60)
    write_done("PASS", 0)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())